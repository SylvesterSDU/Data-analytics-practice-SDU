"""Build site/data/schemas.json: a short description of every dataset for the AI tutor.

Generated from the real files so the tutor uses the actual sheet, table and column
names and data types instead of guessing. Run:  python3 tools/schemas.py
"""
import csv
import json
import sqlite3
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from qbank.common import FILES, ROOT, db_path  # noqa: E402

QUIRKS = {
    "pcards_2013_2014.zip": "TransactionDate and PostedDate are TEXT like '7/26/2014 0:00:00' (not ISO dates). Year and Month are integer columns. FullName = first initial + last name. Negative Amount = return.",
    "Analytics_mindset_case_studies_ETL_Case1.csv": "Comma-separated. 'Apple, Inc.' contains an unquoted comma, so those rows have 9 fields. Cash uses mixed formats (e.g. '20500 M', '900 K'); negatives sometimes in parentheses.",
    "Analytics_mindset_case_studies_IntegrateCo_jobcodes.csv": "Pipe-delimited (|). Job_Code is text with leading zeros (001-012).",
    "Analytics_mindset_case_studies_IntegrateCo_locations.csv": "Pipe-delimited (|). Employee_Name is 'First Last'.",
    "Analytics_mindset_case_studies_IntegrateCo_2015payroll.xlsx": "Job = 2-digit year + optional H/S + 3-digit job code + optional location code (B, SG, R). Gross pay = Reg_Pay + OT_Pay.",
    "Analytics_mindset_case_studies_IntegrateCo_2016payroll.xlsx": "Project uses the same format as Job in 2015. Gross_Wages = regular + overtime pay. Some rows have no Employee_Number.",
    "Analytics_mindset_case_studies_ETL_Case4_Excel.xlsx": "PreparerID is unique only together with BusinessUnitID. Two GLAccountNumber values are text (70300-HP, 70300-FS).",
    "movies.zip": "Contains movies.db. There are two people named Kevin Bacon.",
}


def show(v):
    s = str(v)
    return s if len(s) <= 40 else s[:37] + "..."


def describe_excel(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = ws.iter_rows(max_row=2, values_only=True)
        header = next(rows, ())
        sample = next(rows, ())
        cols = [f"{h} (e.g. {show(s)})" for h, s in zip(header, sample) if h is not None]
        if cols:
            out.append(f"Sheet '{ws.title}' ({ws.max_row - 1} data rows): " + ", ".join(cols))
    return "\n".join(out)


def describe_csv(path):
    text = path.read_text(encoding="utf-8-sig")
    delim = "|" if text.splitlines()[0].count("|") > text.splitlines()[0].count(",") else ","
    rows = list(csv.reader(text.splitlines(), delimiter=delim))
    cols = [f"{h} (e.g. {show(s)})" for h, s in zip(rows[0], rows[1])]
    return f"Delimiter '{delim}', {len(rows) - 1} data rows. Columns: " + ", ".join(cols)


def describe_db(name):
    con = sqlite3.connect(f"file:{db_path(name)}?mode=ro", uri=True)
    out = []
    for (t,) in con.execute("SELECT name FROM sqlite_master WHERE type = 'table' ORDER BY name"):
        n = con.execute(f'SELECT COUNT(*) FROM "{t}"').fetchone()[0]
        cols = [f"{c[1]} {c[2] or ''}".strip() for c in con.execute(f'PRAGMA table_info("{t}")')]
        out.append(f"Table {t} ({n:,} rows): " + ", ".join(cols))
    return "\n".join(out)


schemas = {}
for f in sorted(p.name for p in FILES.iterdir()):
    path = FILES / f
    if f.endswith((".xlsx", ".xlsm")):
        desc = describe_excel(path)
    elif f.endswith(".csv"):
        desc = describe_csv(path)
    else:
        desc = describe_db(f)
    if f in QUIRKS:
        desc += "\nNote: " + QUIRKS[f]
    schemas[f] = desc

(ROOT / "site" / "data" / "schemas.json").write_text(json.dumps(schemas, indent=1))
print(f"{len(schemas)} dataset descriptions, {sum(len(v) for v in schemas.values()):,} characters")
